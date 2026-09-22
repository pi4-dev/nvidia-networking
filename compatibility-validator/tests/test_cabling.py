"""Installation planning regressions using real catalog data and bounded exports."""

import copy
import io
import re
import unittest
from zipfile import ZipFile

from openpyxl import load_workbook
from pydantic import ValidationError

from app.cabling import build_cabling
from app.cabling_exports import cabling_pdf, cabling_xlsx
from app.projects import connection_csv, validate_project
from app.topology_models import CablingOptions, ProjectRequest
from test_topology import SNAPSHOT, cable_breakout, optical_fixture, point_connection, project_of


def report(project, snapshot=SNAPSHOT):
    return {**build_cabling(snapshot, project, validate_project(snapshot, project)),
            'application_version': '0.05-dev', 'evaluated_at': '2026-09-22T12:00:00+00:00'}


def installed_project():
    project = project_of(cable_breakout())
    project.cabling = CablingOptions.model_validate(dict(
        locations=[dict(instance_id='switch-01', rack='Poznań - szafa A', rack_u=12),
                   dict(instance_id='cx8-1', rack='Szafa B', rack_u=20), dict(instance_id='cx8-2', rack='Szafa C', rack_u=21)],
        port_labels=[dict(instance_id=s.instance_id, port_group_id=s.port_group_id, port_number=s.port_number, label=label)
            for s, label in [(project.breakouts[0].head, '1/1'), (project.breakouts[0].branches[0].selection, 'NIC 1'), (project.breakouts[0].branches[1].selection, 'NIC 2')]],
        cables=[dict(entry_id='fanout-1', cable_id='IB-A001', progress=[])]))
    initial = report(project)
    data = project.model_dump()
    data['cabling']['cables'][0]['progress'] = [dict(branch_id=row['branch_id'], installed=True,
        checked=row['branch_id'] == 'branch-1', definition=row['definition'], notes='Odbiór łącza') for row in initial['rows']]
    return ProjectRequest.model_validate(data)


class CablingTests(unittest.TestCase):
    def test_legacy_projects_get_stable_cable_ids_and_one_shared_head_label(self):
        project = project_of(cable_breakout(), connections=[point_connection()])
        data = report(project)
        self.assertEqual(data['summary'], dict(cables=2, legs=3, labels=5, installed=0, checked=0, stale_confirmations=0))
        head = [v for v in data['labels'] if v['end'] == 'H']
        self.assertEqual(len(head), 1)
        self.assertEqual(len(head[0]['peers']), 2)
        self.assertEqual({r['cable_id'] for r in data['rows']}, {'C-link-1', 'C-fanout-1'})
        self.assertTrue(data['provisional'])
        self.assertTrue(any('port label' in x for x in data['issues']))

    def test_progress_round_trip_never_upgrades_compatibility(self):
        project = installed_project()
        restored = ProjectRequest.model_validate_json(project.model_dump_json())
        data = report(restored)
        self.assertEqual(data['summary']['installed'], 2)
        self.assertEqual(data['summary']['checked'], 1)
        self.assertEqual(data['summary']['stale_confirmations'], 0)
        self.assertEqual(data['status'], 'unknown')
        self.assertTrue(data['provisional'])
        self.assertEqual(data['issues'], [])
        self.assertEqual(data['rows'][0]['source']['rack'], 'Poznań - szafa A')
        self.assertEqual(data['rows'][0]['source']['port_label'], '1/1')

    def test_changed_definition_requires_fresh_installation_confirmation(self):
        for field in ('port', 'rack', 'cable_id', 'pn', 'length'):
            project = installed_project()
            if field == 'port': project.breakouts[0].head.port_number = 2
            if field == 'rack': project.cabling.locations[0].rack = 'Different rack'
            if field == 'cable_id': project.cabling.cables[0].cable_id = 'IB-A002'
            if field == 'pn': project.breakouts[0].head.part_number = 'UNKNOWN-PN'
            if field == 'length': project.breakouts[0].length_m = 5.
            data = report(project)
            with self.subTest(field=field):
                self.assertEqual(data['summary']['installed'], 0)
                self.assertEqual(data['summary']['checked'], 0)
                self.assertEqual(data['summary']['stale_confirmations'], 2)

    def test_reordering_branches_or_editing_notes_preserves_progress(self):
        project = installed_project()
        before = report(project)
        project.breakouts[0].branches.reverse()
        project.cabling.cables[0].progress[0].notes = 'New note'
        after = report(project)
        self.assertEqual(before['summary'], after['summary'])
        self.assertEqual([r['definition'] for r in before['rows']], [r['definition'] for r in after['rows']])

    def test_cable_id_collision_includes_generated_ids_case_insensitively(self):
        project = project_of(cable_breakout(), connections=[point_connection()]).model_dump()
        project['cabling']['cables'] = [dict(entry_id='fanout-1', cable_id='c-LINK-1')]
        with self.assertRaisesRegex(ValidationError, 'duplicate cable ID'):
            ProjectRequest.model_validate(project)

    def test_orphaned_and_invalid_metadata_is_rejected(self):
        cases = [dict(locations=[dict(instance_id='missing', rack='A')]),
                 dict(port_labels=[dict(instance_id='switch-01', port_group_id='ndr', port_number=2, label='2')]),
                 dict(cables=[dict(entry_id='missing')]),
                 dict(cables=[dict(entry_id='fanout-1', progress=[dict(branch_id='missing')])]),
                 dict(locations=[dict(instance_id='switch-01', rack='bad\x00label')]),
                 dict(cables=[dict(entry_id='fanout-1', progress=[dict(branch_id='branch-1', checked=True)])]),
                 dict(cables=[dict(entry_id='fanout-1', progress=[dict(branch_id='branch-1', installed=True)])])]
        for value in cases:
            data = project_of(cable_breakout()).model_dump(); data['cabling'] = value
            with self.subTest(value=value), self.assertRaises(ValidationError): ProjectRequest.model_validate(data)

    def test_optical_breakout_labels_describe_harness_and_keep_module_pns_separate(self):
        snapshot, breakout = optical_fixture()
        data = report(project_of(breakout), snapshot)
        self.assertEqual(data['summary']['labels'], 5)
        self.assertTrue(all(r['part_number'] == 'FIXTURE-HARNESS' for r in data['rows']))
        self.assertEqual(data['rows'][0]['source']['module_pn'], '980-9I16Y-00W000')
        self.assertEqual(data['rows'][0]['destination']['module_pn'], '980-9I042-00C000')
        self.assertEqual(data['rows'][0]['head_optical_lanes'], [1])

    def test_missing_declared_length_uses_documented_exact_sku(self):
        project = project_of(connections=[point_connection()]); project.connections[0].length_m = None
        data = report(project)
        self.assertEqual(data['rows'][0]['length_m'], 3.)
        self.assertEqual(data['rows'][0]['documented_length_m'], 3.)

    def test_changed_documented_length_invalidates_confirmation(self):
        project = project_of(connections=[point_connection()]); project.connections[0].length_m = None
        initial = report(project)['rows'][0]
        project.cabling = CablingOptions(cables=[dict(entry_id=initial['entry_id'], progress=[dict(
            installed=True, checked=True, definition=initial['definition'])])])
        self.assertEqual(report(project)['summary']['checked'], 1)
        snapshot = copy.deepcopy(SNAPSHOT)
        for product in snapshot['interconnects']:
            for sku in product['skus']:
                if sku['part_number'] == initial['part_number']: sku['length_m'] = 5.
        changed = report(project, snapshot)
        self.assertEqual(changed['rows'][0]['length_m'], 5.)
        self.assertEqual(changed['summary']['checked'], 0)
        self.assertEqual(changed['summary']['stale_confirmations'], 1)

    def test_csv_cannot_silently_drop_installation_metadata(self):
        project = project_of(connections=[point_connection()])
        project.cabling = CablingOptions(locations=[dict(instance_id='switch-point', rack='Rack A')])
        with self.assertRaisesRegex(ValueError, 'cabling locations'): connection_csv(project)

    def test_xlsx_preserves_literal_values_numbers_booleans_and_all_labels(self):
        data = report(installed_project()); data['name'] = '=HYPERLINK("https://example.com")'
        data['rows'][0]['notes'] = '=1+1'
        contents = cabling_xlsx(data)
        book = load_workbook(io.BytesIO(contents))
        self.assertEqual(book.sheetnames, ['Cabling plan', 'End labels'])
        sheet = book['Cabling plan']
        self.assertEqual(sheet['A1'].data_type, 's')
        self.assertEqual(sheet['K6'].value, '=1+1')
        self.assertEqual(sheet['K6'].data_type, 's')
        self.assertEqual(sheet['F6'].value, 3)
        self.assertIs(sheet['H6'].value, True)
        self.assertIs(sheet['I6'].value, True)
        self.assertEqual(sheet.freeze_panes, 'C6')
        self.assertEqual(sheet.auto_filter.ref, 'A5:K7')
        self.assertEqual(book['End labels'].max_row, 8)
        self.assertIn('Poznań', sheet['C6'].value)
        with ZipFile(io.BytesIO(contents)) as archive:
            self.assertFalse(any(b'<f>' in archive.read(n) for n in archive.namelist() if n.startswith('xl/worksheets/')))

    def test_pdf_exports_embed_font_and_paginate_without_dropping_long_rows(self):
        data = report(installed_project()); data['name'] = 'Łącza - Żółć <projekt>'
        data['rows'][0]['notes'] = 'Długi opis montażu & odbioru. ' * 16
        data['rows'][0]['cable_id'] = 'C-' + 'LONG-ID-' * 8
        data['rows'] = data['rows'] * 16
        data['labels'] = data['labels'] * 12
        for labels in (False, True):
            contents = cabling_pdf(data, labels=labels)
            self.assertTrue(contents.startswith(b'%PDF-'))
            self.assertTrue(contents.rstrip().endswith(b'%%EOF'))
            self.assertIn(b'/FontFile2', contents)
            self.assertGreater(max(map(int, re.findall(rb'/Count (\d+)', contents))), 1)


if __name__ == '__main__':
    unittest.main()
