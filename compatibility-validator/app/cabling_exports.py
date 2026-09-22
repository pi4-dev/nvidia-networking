"""Self-contained installation documents; no network or temporary files at runtime."""

import io
import math
import re
from html import escape
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, LongTable

from .cabling import endpoint_text

FONT = "CablingSans"
pdfmetrics.registerFont(TTFont(FONT, str(Path(__file__).parent / "fonts" / "DejaVuSans.ttf")))
INK, MUTED, ACCENT = "173047", "486276", "DDEBF3"
CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]")


def clean(value):
    return CONTROL.sub("\ufffd", str(value if value is not None else ""))


def branch_text(row):
    if row["branch_id"] is None:
        return row["entry_id"]
    lines = [row["entry_id"], row["branch_id"], "Termination " + str(row["termination"]),
             "Head links: " + ",".join(map(str, row["head_links"]))]
    if row["head_optical_lanes"]:
        lines.append("Optical pairs: " + ",".join(map(str, row["head_optical_lanes"])) + " -> " + ",".join(map(str, row["branch_optical_lanes"])))
    return "\n".join(lines)


def label_peer_text(label):
    if not label["peers"]:
        return "Unassigned - review the complete cabling plan"
    if len(label["peers"]) == 1:
        return endpoint_text(label["peers"][0])
    return f"{len(label['peers'])} branches - see the cabling plan for individual ports"


def _paragraph(value, size=8, color=INK):
    return Paragraph(escape(clean(value)).replace("\n", "<br/>"), ParagraphStyle(
        "cabling", fontName=FONT, fontSize=size, leading=size * 1.35,
        textColor=colors.HexColor("#" + color), alignment=TA_LEFT, splitLongWords=True, spaceAfter=0))


def _metadata(report):
    return (f"Catalog {report['revision']} | Application {report['application_version']} | {report['evaluated_at']}\n"
            f"Compatibility: {report['status']} | {'PROVISIONAL' if report['provisional'] else 'Validated plan'} | "
            f"{report['summary']['cables']} assemblies / {report['summary']['legs']} legs / {report['summary']['labels']} labels")


def cabling_pdf(report, labels=False):
    stream = io.BytesIO()
    page = A4 if labels else landscape(A4)
    doc = SimpleDocTemplate(stream, pagesize=page, leftMargin=24, rightMargin=24, topMargin=28, bottomMargin=35,
                            title="Cable labels" if labels else "Cabling plan", author="NVIDIA Networking Compatibility Validator")
    width = page[0] - 48
    story = [_paragraph("CABLE LABELS" if labels else "CABLING PLAN", 18), Spacer(1, 8),
             _paragraph(report["name"], 12), Spacer(1, 6), _paragraph(_metadata(report), 8, MUTED), Spacer(1, 12)]
    if labels:
        story += [_paragraph("Print on A4 at 100% scale. Cut along the borders; cards grow to fit their text. Shared breakout heads appear once per physical connector.", 8), Spacer(1, 10)]
        # Independent rows allow variable-size cards without losing long identifiers.
        for offset in range(0, len(report["labels"]), 2):
            cells = []
            for label in report["labels"][offset:offset + 2]:
                cells.append([_paragraph(label["label_id"], 11), Spacer(1, 5),
                    _paragraph("LOCAL\n" + endpoint_text(label["local"]), 8), Spacer(1, 5),
                    _paragraph("REMOTE\n" + label_peer_text(label), 8), Spacer(1, 5),
                    _paragraph("Cable PN: " + (label["part_number"] or "UNKNOWN"), 8),
                    _paragraph("PROVISIONAL" if report["provisional"] else "Validated plan", 7, MUTED)])
            if len(cells) == 1:
                cells.append("")
            table = Table([cells], colWidths=[width / 2] * 2, hAlign="LEFT")
            table.setStyle(TableStyle([("BOX", (0, 0), (0, 0), .5, colors.HexColor("#" + MUTED)),
                *([("BOX", (1, 0), (1, 0), .5, colors.HexColor("#" + MUTED))] if len(report["labels"][offset:offset + 2]) == 2 else []),
                ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 10),
                ("RIGHTPADDING", (0, 0), (-1, -1), 10), ("TOPPADDING", (0, 0), (-1, -1), 10), ("BOTTOMPADDING", (0, 0), (-1, -1), 10)]))
            story += [table, Spacer(1, 10)]
    else:
        story += [_paragraph("Installed/checked are recorded by the installer and do not change compatibility. Blank locations or unconfirmed port markings require review.", 8), Spacer(1, 10)]
        headers = ["Cable / branch", "Source", "Destination", "Cable PN / length", "Installation"]
        data = [[_paragraph(h, 8) for h in headers]]
        for row in report["rows"]:
            length = f"{row['length_m']:g} m" if row["length_m"] is not None else "Length unknown"
            state = f"Installed: {'YES' if row['installed'] else 'NO'}\nChecked: {'YES' if row['checked'] else 'NO'}\n{row['validation_status']}"
            if row["progress_stale"]:
                state += "\nPrevious confirmation is stale"
            if row["notes"]:
                state += "\n" + row["notes"]
            data.append([_paragraph(row["cable_id"] + "\n" + branch_text(row)),
                _paragraph(endpoint_text(row["source"], model=True)), _paragraph(endpoint_text(row["destination"], model=True)),
                _paragraph((row["part_number"] or "PN unknown") + "\n" + length + "\n" + row["fabric"]), _paragraph(state)])
        table = LongTable(data, colWidths=[width * p for p in (.17, .245, .245, .15, .19)], repeatRows=1, hAlign="LEFT")
        table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + ACCENT)),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F7FA")]),
            ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 7),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 7), ("LINEBELOW", (0, 0), (-1, 0), .6, colors.HexColor("#" + MUTED))]))
        story.append(table)
        if report["issues"]:
            story += [Spacer(1, 14), _paragraph("Installation details to complete", 11), Spacer(1, 6)]
            for issue in report["issues"]:
                story += [_paragraph(issue), Spacer(1, 3)]

    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont(FONT, 7)
        canvas.setFillColor(colors.HexColor("#" + MUTED))
        canvas.drawString(24, 19, "PROVISIONAL | " + report["revision"] if report["provisional"] else "Catalog " + report["revision"])
        canvas.drawRightString(page[0] - 24, 19, f"Page {document.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return stream.getvalue()


def cabling_xlsx(report):
    book = Workbook()
    plan = book.active; plan.title = "Cabling plan"
    labels = book.create_sheet("End labels")
    headers = ["Cable ID", "Entry / branch", "Source", "Destination", "Cable PN", "Length (m)", "Fabric", "Installed", "Checked", "Compatibility", "Notes"]
    plan_rows = []
    for row in report["rows"]:
        notes = ("Previous confirmation is stale. " if row["progress_stale"] else "") + row["notes"]
        plan_rows.append([row["cable_id"], branch_text(row), endpoint_text(row["source"], model=True),
            endpoint_text(row["destination"], model=True), row["part_number"] or "PN unknown", row["length_m"],
            row["fabric"], row["installed"], row["checked"], row["validation_status"], notes])
    label_rows = [[v["label_id"], v["cable_id"], v["end"], endpoint_text(v["local"]), label_peer_text(v), v["part_number"] or "PN unknown"] for v in report["labels"]]

    for sheet, titles, data, widths in [(plan, headers, plan_rows, [24, 24, 42, 42, 26, 12, 10, 12, 12, 16, 48]),
        (labels, ["Label ID", "Cable ID", "End", "Local endpoint", "Remote endpoint(s)", "Cable PN"], label_rows, [34, 24, 24, 45, 45, 28])]:
        sheet.append([report["name"]]); sheet.append([_metadata(report)])
        sheet.append(["Installer declarations are separate from compatibility. Update progress in the app; use project JSON to retain editable installation data."])
        sheet.append([]); sheet.append(titles)
        for row in data:
            sheet.append(row)
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    # Force literal text even for imported values beginning with '='.
                    cell.value = clean(cell.value); cell.data_type = "s"
                cell.font = Font(name="Calibri", size=11, color=INK)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for r in (1, 2, 3):
            sheet.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(titles))
        sheet["A1"].font = Font(name="Calibri", size=18, bold=True, color=INK)
        sheet.row_dimensions[1].height = 45
        sheet.row_dimensions[2].height = 45
        sheet.row_dimensions[3].height = 36
        for cell in sheet[5]:
            cell.fill = PatternFill("solid", fgColor=INK)
            cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        sheet.row_dimensions[5].height = 32
        for i, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(i)].width = width
        for r in range(6, sheet.max_row + 1):
            lines = max(sum(max(1, math.ceil(len(line) / max(1, widths[i] - 3))) for line in str(c.value or "").split("\n")) for i, c in enumerate(sheet[r]))
            sheet.row_dimensions[r].height = max(36, min(409, lines * 16 + 8))
            if r % 2 == 0:
                for cell in sheet[r]:
                    cell.fill = PatternFill("solid", fgColor="F0F5F8")
        sheet.freeze_panes = "C6"
        sheet.auto_filter.ref = f"A5:{get_column_letter(len(titles))}{sheet.max_row}"
        sheet.print_title_rows = "1:5"
        sheet.sheet_view.showGridLines = False
        sheet.page_setup.orientation = "landscape"
        sheet.print_options.horizontalCentered = True
    for cell in plan["F"][5:]:
        cell.number_format = "0.0##"
    if report["issues"]:
        # Keep completion notes below the filterable plan, outside its data range.
        start = plan.max_row + 3
        plan.cell(start, 1, "Installation details to complete")
        for n, issue in enumerate(report["issues"], start + 1):
            cell = plan.cell(n, 1, clean(issue)); cell.data_type = "s"
            plan.merge_cells(start_row=n, start_column=1, end_row=n, end_column=len(headers))
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            plan.row_dimensions[n].height = 30
    out = io.BytesIO(); book.save(out)
    return out.getvalue()
